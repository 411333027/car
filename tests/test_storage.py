import os

import pytest
from openpyxl import load_workbook

from gas_agent import storage
from gas_agent.storage import RecordStore, StorageBusy


def make_store(tmp_path, sleeps=None, times=3):
    return RecordStore(
        tmp_path / "records.xlsx", tmp_path / "backup",
        retry_seconds=0, retry_times=times, sleep=(sleeps.append if sleeps is not None else (lambda s: None)),
    )


def row(uid="123456789012345678", **kw):
    base = {"time": "2026-10-26 10:00:00", "user_id": uid, "username": "grace", "type": "加油",
            "odometer": 12000, "liters": 30.5, "cost": 900, "full": "是", "status": "完成"}
    return {**base, **kw}


def test_appends_rows_without_overwriting(tmp_path):
    store = make_store(tmp_path)
    ids = [store.append_row(row(odometer=12000 + i)) for i in range(3)]
    assert ids == [1, 2, 3]
    # 重新建立 store（模擬重啟）後繼續往下加
    ids2 = make_store(tmp_path).append_row(row(odometer=13000))
    assert ids2 == 4
    ws = load_workbook(tmp_path / "records.xlsx").active
    assert ws.max_row == 5  # 標題 + 4 筆
    assert [ws.cell(r, 1).value for r in range(2, 6)] == [1, 2, 3, 4]


def test_user_id_stored_as_text_keeps_all_digits(tmp_path):
    store = make_store(tmp_path)
    store.append_row(row("123456789012345678"))
    ws = load_workbook(tmp_path / "records.xlsx").active
    assert ws.cell(2, 3).value == "123456789012345678"
    assert store.rows_for_user("123456789012345678")[0]["user_id"] == "123456789012345678"


def test_rows_are_isolated_per_user(tmp_path):
    store = make_store(tmp_path)
    store.append_row(row("111", odometer=100))
    store.append_row(row("222", odometer=200))
    store.append_row(row("111", odometer=150))
    assert [r["odometer"] for r in store.rows_for_user("111")] == [100, 150]
    assert [r["odometer"] for r in store.rows_for_user("222")] == [200]
    assert store.rows_for_user("333") == []


@pytest.mark.parametrize("evil", ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1)", '=HYPERLINK("http://x","y")'])
def test_formula_injection_is_neutralized(tmp_path, evil):
    store = make_store(tmp_path)
    store.append_row(row(username=evil, note=evil, maint_item=evil))
    ws = load_workbook(tmp_path / "records.xlsx").active
    for col in (4, 13, 15):  # 使用者名稱、保養項目、說明
        cell = ws.cell(2, col)
        assert cell.data_type == "s", f"{evil!r} 被存成 {cell.data_type}"
        assert cell.value.startswith("'")
    assert ws.cell(2, 6).value == 12000  # 數字欄位仍是數字


def test_update_row(tmp_path):
    store = make_store(tmp_path)
    store.append_row(row())
    store.update_row(1, {"odometer": 12400, "status": "已修正"})
    r = store.rows_for_user("123456789012345678")[0]
    assert r["odometer"] == 12400 and r["status"] == "已修正"


def test_delete_row(tmp_path):
    store = make_store(tmp_path)
    store.append_row(row("1"))
    store.append_row(row("2"))
    store.delete_row(2)
    assert store.rows_for_user("2") == []
    assert len(store.rows_for_user("1")) == 1


def test_locked_file_retries_then_succeeds(tmp_path, monkeypatch, caplog):
    store = make_store(tmp_path)
    store.append_row(row())
    real_replace = os.replace
    calls = {"n": 0}

    def flaky_replace(src, dst):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError("file is open in Excel")
        return real_replace(src, dst)

    monkeypatch.setattr(storage.os, "replace", flaky_replace)
    sleeps = []
    store2 = make_store(tmp_path, sleeps, times=5)
    with caplog.at_level("WARNING"):
        store2.append_row(row(odometer=12500))
    assert calls["n"] == 3 and len(sleeps) == 2
    assert "占用" in caplog.text
    assert len(store2.rows_for_user("123456789012345678")) == 2  # 沒有重複寫入


def test_locked_file_gives_up_with_storage_busy(tmp_path, monkeypatch):
    store = make_store(tmp_path, times=3)
    store.append_row(row())

    def always_locked(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(storage.os, "replace", always_locked)
    with pytest.raises(StorageBusy):
        store.append_row(row(odometer=1))
    monkeypatch.undo()
    assert len(store.rows_for_user("123456789012345678")) == 1  # 原檔沒壞、沒有半筆
    assert not list(tmp_path.glob("*.tmp"))
