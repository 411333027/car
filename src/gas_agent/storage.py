"""Excel 紀錄檔（DATA_DIR/records.xlsx）：一直往下加、依 user_id 分人、檔案被占用時重試。

這個檔案有一段防護：
- 防護 4：寫進 Excel 的文字若以 = + - @ 開頭，先處理成純文字，避免打開 Excel 時被當公式
"""

import logging
import os
import shutil
import time
from collections.abc import Callable
from datetime import date
from pathlib import Path
from typing import TypeVar

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from gas_agent.sanitize import neutralize_formula

log = logging.getLogger(__name__)
T = TypeVar("T")

SHEET = "紀錄"
# (內部名稱, Excel 標題, 欄寬)
COLUMNS = [
    ("id", "紀錄編號", 9),
    ("time", "日期時間", 20),
    ("user_id", "使用者ID", 22),
    ("username", "使用者名稱", 16),
    ("type", "紀錄類型", 12),
    ("odometer", "當前總里程(km)", 15),
    ("trip", "本次行駛距離(km)", 17),
    ("liters", "加油量(L)", 11),
    ("cost", "花費金額(元)", 13),
    ("full", "是否加滿", 10),
    ("kml", "平均油耗(km/L)", 15),
    ("cost_per_km", "每公里成本(元/km)", 18),
    ("maint_item", "保養項目", 22),
    ("status", "狀態", 11),
    ("note", "說明", 40),
    ("summary", "個人累積總統計", 60),
    ("photo", "照片存放位置", 50),
]
KEYS = [c[0] for c in COLUMNS]
KEY_TO_COL = {key: i + 1 for i, key in enumerate(KEYS)}
TEXT_KEYS = ("time", "username", "type", "full", "maint_item", "status", "note", "summary", "photo")


class StorageBusy(Exception):
    """紀錄檔一直被占用（通常是 Excel 開著），重試用完還是寫不進去。"""


class StorageError(Exception):
    """紀錄檔壞掉或其他無法處理的問題。"""


def _put(cell, value) -> None:
    """寫一個儲存格。文字一律走防護 4。"""
    cell.value = value  # 沒有防護 4 時的預設：openpyxl 會把 "=..." 當成公式

    # ===== 防護 4：Excel 公式注入處理 開始 =====
    # 使用者名稱、保養項目、AI 的說明都是「別人給的字」。以 = + - @ 開頭時，
    # 前面加單引號並明確設成「文字」格式，Excel 就只會把它當字，不會執行。
    # 數字欄位（里程、金額…）不是字串，不受影響。
    if isinstance(value, str):
        cell.value = neutralize_formula(value)
        cell.data_type = "s"
        cell.number_format = "@"
    # ===== 防護 4 結束 =====


def _as_user_id(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class RecordStore:
    def __init__(
        self,
        path: Path,
        backup_dir: Path,
        retry_seconds: float = 5,
        retry_times: int = 12,
        sleep: Callable[[float], None] = time.sleep,
    ):
        self.path = path
        self.backup_dir = backup_dir
        self.retry_seconds = retry_seconds
        self.retry_times = retry_times
        self._sleep = sleep
        self._backed_up_on: date | None = None

    # ---------- 內部：載入、重試、原子寫入 ----------
    def _load(self) -> tuple[Workbook, object]:
        if self.path.exists():
            wb = load_workbook(self.path)
            ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
            return wb, ws
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        for col, (_key, header, width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"
        return wb, ws

    def _save(self, wb: Workbook) -> None:
        """先存成暫存檔，成功後再整份替換原檔；中途出事不會留下壞掉的一半。"""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_name(self.path.name + ".tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        finally:
            tmp.unlink(missing_ok=True)

    def _retrying(self, op: Callable[[], T], what: str) -> T:
        for attempt in range(1, self.retry_times + 1):
            try:
                return op()
            except PermissionError:
                log.warning(
                    "⚠️ 紀錄檔 %s 正被占用（Excel 可能開著），無法%s。請關閉 Excel。（第 %d/%d 次）",
                    self.path.name, what, attempt, self.retry_times,
                )
                if attempt == self.retry_times:
                    break
                self._sleep(self.retry_seconds)
            except StorageError:
                raise
            except Exception as exc:  # 檔案壞掉等：重試也沒用
                raise StorageError(f"{what}失敗：{exc}") from exc
        raise StorageBusy(f"紀錄檔一直被占用，{what}失敗")

    def _backup_once_a_day(self) -> None:
        today = date.today()
        if self._backed_up_on == today or not self.path.exists():
            return
        try:
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            target = self.backup_dir / f"records-{today:%Y%m%d}.xlsx"
            if not target.exists():
                shutil.copyfile(self.path, target)
            for old in sorted(self.backup_dir.glob("records-*.xlsx"))[:-7]:
                old.unlink(missing_ok=True)
            self._backed_up_on = today
        except OSError as exc:
            log.warning("備份失敗（不影響記帳）：%s", exc)

    @staticmethod
    def _row_to_dict(ws, row_idx: int) -> dict:
        row = {key: ws.cell(row=row_idx, column=col).value for key, col in KEY_TO_COL.items()}
        row["user_id"] = _as_user_id(row["user_id"])
        for key in TEXT_KEYS:
            row[key] = "" if row[key] is None else str(row[key])
        return row

    @staticmethod
    def _find_row_idx(ws, record_id: int) -> int | None:
        for idx in range(2, ws.max_row + 1):
            if ws.cell(row=idx, column=KEY_TO_COL["id"]).value == record_id:
                return idx
        return None

    # ---------- 對外：讀 ----------
    def rows_for_user(self, user_id: str) -> list[dict]:
        """只回傳這個 user_id 的紀錄（依時間順序）。沒有「讀全部人」的函式，分人隔離就靠這一關。"""
        wanted = str(int(user_id))

        def op() -> list[dict]:
            if not self.path.exists():
                return []
            wb = load_workbook(self.path)
            try:
                ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
                rows = [self._row_to_dict(ws, i) for i in range(2, ws.max_row + 1)]
            finally:
                wb.close()
            return [r for r in rows if r["id"] is not None and r["user_id"] == wanted]

        return self._retrying(op, "讀取紀錄")

    # ---------- 對外：寫 ----------
    def append_row(self, row: dict) -> int:
        """在最後面新增一列（絕不覆蓋），回傳新的紀錄編號。"""

        def op() -> int:
            self._backup_once_a_day()
            wb, ws = self._load()
            ids = [
                ws.cell(row=i, column=1).value
                for i in range(2, ws.max_row + 1)
                if isinstance(ws.cell(row=i, column=1).value, int)
            ]
            new_id = (max(ids) if ids else 0) + 1
            target = ws.max_row + 1
            full_row = {**row, "id": new_id}
            for key, col in KEY_TO_COL.items():
                _put(ws.cell(row=target, column=col), full_row.get(key))
            self._save(wb)
            return new_id

        return self._retrying(op, "寫入紀錄")

    def update_row(self, record_id: int, updates: dict) -> None:
        def op() -> None:
            self._backup_once_a_day()
            wb, ws = self._load()
            idx = self._find_row_idx(ws, record_id)
            if idx is None:
                raise StorageError(f"找不到紀錄編號 {record_id}")
            for key, value in updates.items():
                _put(ws.cell(row=idx, column=KEY_TO_COL[key]), value)
            self._save(wb)

        self._retrying(op, "修改紀錄")

    def delete_row(self, record_id: int) -> None:
        """退回一筆剛寫入的紀錄（寫入後發現後續步驟失敗時使用）。"""

        def op() -> None:
            wb, ws = self._load()
            idx = self._find_row_idx(ws, record_id)
            if idx is not None:
                ws.delete_rows(idx)
                self._save(wb)

        self._retrying(op, "退回紀錄")
